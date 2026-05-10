from __future__ import annotations

from copy import copy
from io import BytesIO
from numbers import Real
import re

from openpyxl.styles import Font
from openpyxl.styles import PatternFill
import pandas as pd


MONTH_PATTERN = re.compile(r"(20\d{2}(0[1-9]|1[0-2]))")
MONTH_NAME_PATTERN = re.compile(
    r"^(januar|februar|mars|april|mai|juni|juli|august|september|oktober|november|desember)( \[\d+\])?$",
    re.IGNORECASE,
)
MONTH_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")
SUMMARY_FILLS = {
    "regnskap": PatternFill(fill_type="solid", fgColor="DDEBF7"),
    "budsjett": PatternFill(fill_type="solid", fgColor="E2F0D9"),
    "avvik": PatternFill(fill_type="solid", fgColor="FCE4EC"),
}
VALUE_COLUMN_KEYWORDS = ("regnskap", "budsjett", "avvik", "forbruk", "prognose")


def build_export_workbook(
    cleaned_df: pd.DataFrame,
    adjustments_df: pd.DataFrame,
    user_notes: str,
    source_name: str,
) -> bytes:
    buffer = BytesIO()

    notes_df = pd.DataFrame(
        [
            {"Felt": "Kilde", "Verdi": source_name},
            {"Felt": "Brukerkommentar", "Verdi": user_notes.strip() or "Ingen kommentar oppgitt."},
        ]
    )

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        cleaned_df.to_excel(writer, sheet_name="Renset data", index=False)
        adjustments_df.to_excel(writer, sheet_name="Endringslogg", index=False)
        notes_df.to_excel(writer, sheet_name="Forklaringer", index=False)

        for sheet_name, dataframe in {
            "Renset data": cleaned_df,
            "Endringslogg": adjustments_df,
            "Forklaringer": notes_df,
        }.items():
            worksheet = writer.sheets[sheet_name]
            if dataframe.empty:
                continue
            for index, column_name in enumerate(dataframe.columns, start=1):
                values = [str(column_name)] + dataframe[column_name].astype(str).tolist()
                worksheet.column_dimensions[worksheet.cell(row=1, column=index).column_letter].width = min(
                    max(len(value) for value in values) + 2,
                    60,
                )

        _apply_month_column_colors(writer.sheets["Renset data"], cleaned_df)
        _apply_currency_number_format(writer.sheets["Renset data"], cleaned_df)
        _apply_subtotal_row_formatting(writer.sheets["Renset data"], cleaned_df)

    buffer.seek(0)
    return buffer.getvalue()


def _apply_month_column_colors(worksheet, cleaned_df: pd.DataFrame) -> None:
    for column_index, column_name in enumerate(cleaned_df.columns, start=1):
        summary_fill = _extract_summary_fill(str(column_name))
        if summary_fill is not None:
            for row_index in range(1, worksheet.max_row + 1):
                worksheet.cell(row=row_index, column=column_index).fill = summary_fill
            continue

        month_value = _extract_month_group(str(column_name))
        if not month_value:
            continue
        for row_index in range(1, worksheet.max_row + 1):
            worksheet.cell(row=row_index, column=column_index).fill = MONTH_FILL


def _extract_month_group(column_name: str) -> str | None:
    month_match = MONTH_PATTERN.search(column_name)
    if month_match:
        return month_match.group(1)

    month_name_match = MONTH_NAME_PATTERN.match(column_name.strip())
    if month_name_match:
        return month_name_match.group(1).lower()

    return None


def _extract_summary_fill(column_name: str) -> PatternFill | None:
    normalized_name = column_name.strip().lower()
    return SUMMARY_FILLS.get(normalized_name)


def _apply_currency_number_format(worksheet, cleaned_df: pd.DataFrame) -> None:
    rightmost_month_column = max(
        (
            column_index
            for column_index, column_name in enumerate(cleaned_df.columns, start=1)
            if _extract_month_group(str(column_name))
        ),
        default=0,
    )
    value_columns = {
        column_index
        for column_index, column_name in enumerate(cleaned_df.columns, start=1)
        if _is_value_column(str(column_name))
        or (
            column_index > rightmost_month_column
            and _is_mostly_numeric_series(cleaned_df.iloc[:, column_index - 1])
        )
    }

    for row_index in range(2, worksheet.max_row + 1):
        for column_index in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            if column_index in value_columns:
                coerced_value = _coerce_numeric_cell_value(cell.value)
                if coerced_value is not None:
                    cell.value = coerced_value
            if isinstance(cell.value, Real) and not isinstance(cell.value, bool):
                cell.number_format = "#,##0"


def _is_value_column(column_name: str) -> bool:
    normalized_name = column_name.strip().lower()
    if _extract_month_group(column_name):
        return True
    if "som total" in normalized_name and any(keyword in normalized_name for keyword in VALUE_COLUMN_KEYWORDS):
        return True
    return bool(re.search(r"\b20\d{2}\b", normalized_name)) and any(
        keyword in normalized_name for keyword in VALUE_COLUMN_KEYWORDS
    )


def _coerce_numeric_cell_value(value: object) -> int | float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Real):
        return int(value) if float(value).is_integer() else float(value)

    text = str(value).replace("\xa0", " ").strip()
    if not text:
        return None

    text = text.replace("kr", "").replace("%", "").strip()
    is_negative = False
    if text.startswith("(") and text.endswith(")"):
        is_negative = True
        text = text[1:-1].strip()
    if text.endswith("-"):
        is_negative = True
        text = text[:-1].strip()
    text = re.sub(r"[^0-9,\.\-]", "", text)
    if not text:
        return None

    if text.count(",") == 1 and text.count(".") >= 1:
        text = text.replace(".", "").replace(",", ".")
    elif text.count(",") == 1 and text.count(".") == 0:
        text = text.replace(",", ".")
    elif text.count(",") > 1 and text.count(".") == 0:
        text = text.replace(",", "")
    else:
        text = text.replace(",", "")

    try:
        numeric_value = float(text)
    except ValueError:
        return None

    if is_negative:
        numeric_value = -abs(numeric_value)
    return int(numeric_value) if numeric_value.is_integer() else numeric_value


def _is_mostly_numeric_series(series: pd.Series, threshold: float = 0.8) -> bool:
    non_empty_values = [value for value in series.tolist() if str(value).strip() not in {"", "None", "nan"}]
    if not non_empty_values:
        return False

    numeric_values = [value for value in non_empty_values if _coerce_numeric_cell_value(value) is not None]
    return (len(numeric_values) / len(non_empty_values)) >= threshold


def _apply_subtotal_row_formatting(worksheet, cleaned_df: pd.DataFrame) -> None:
    highlighted_labels = set(cleaned_df.attrs.get("subtotal_labels", []))
    total_label = cleaned_df.attrs.get("total_label")
    if total_label:
        highlighted_labels.add(total_label)
    if not highlighted_labels or worksheet.max_column == 0:
        return

    for row_index in range(2, worksheet.max_row + 1):
        first_value = worksheet.cell(row=row_index, column=1).value
        if str(first_value).strip() not in highlighted_labels:
            continue

        for column_index in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            updated_font = copy(cell.font) if cell.font else Font()
            updated_font.bold = True
            cell.font = updated_font
