from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
import pandas as pd

from kommune_cleaner.exporting import build_export_workbook


def test_build_export_workbook_creates_expected_sheets():
    cleaned_df = pd.DataFrame(
        [
            {
                "Ansvar": "100",
                "januar": 12.5,
                "februar": 14.5,
                "Regnskap": 15.5,
                "Budsjett": 16.5,
                "Avvik": 17.5,
            }
        ]
    )
    adjustments_df = pd.DataFrame(
        [{"Kategori": "Forward fill", "Beskrivelse": "Fylte ned Ansvar.", "Verdi": 3}]
    )

    workbook_bytes = build_export_workbook(
        cleaned_df=cleaned_df,
        adjustments_df=adjustments_df,
        user_notes="Testkommentar",
        source_name="pivotgrid.xlsx - Ark1",
    )

    workbook = load_workbook(BytesIO(workbook_bytes))

    assert workbook.sheetnames == ["Renset data", "Endringslogg", "Forklaringer"]
    assert workbook["Renset data"]["A2"].value == "100"
    assert workbook["Endringslogg"]["A2"].value == "Forward fill"
    assert workbook["Forklaringer"]["B3"].value == "Testkommentar"
    assert workbook["Renset data"]["B2"].number_format == "#,##0"
    assert workbook["Renset data"]["C2"].number_format == "#,##0"
    assert workbook["Renset data"]["D2"].number_format == "#,##0"
    assert workbook["Renset data"]["E2"].number_format == "#,##0"
    assert workbook["Renset data"]["F2"].number_format == "#,##0"
    assert workbook["Renset data"]["B1"].fill.fgColor.rgb == "00F2F2F2"
    assert workbook["Renset data"]["C1"].fill.fgColor.rgb == "00D9D9D9"
    assert workbook["Renset data"]["D1"].fill.fgColor.rgb == "00DDEBF7"
    assert workbook["Renset data"]["E1"].fill.fgColor.rgb == "00E2F0D9"
    assert workbook["Renset data"]["F1"].fill.fgColor.rgb == "00FCE4EC"


def test_build_export_workbook_converts_summary_text_cells_to_numeric_excel_cells():
    cleaned_df = pd.DataFrame(
        [
            {
                "Ansvar": "100",
                "januar": "1 250,00",
                "Regnskap": "12 500",
                "Budsjett": "13 000",
                "Avvik": "500",
            }
        ]
    )

    workbook_bytes = build_export_workbook(
        cleaned_df=cleaned_df,
        adjustments_df=pd.DataFrame(columns=["Kategori", "Beskrivelse", "Verdi"]),
        user_notes="",
        source_name="pivotgrid.xlsx - Ark1",
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    worksheet = workbook["Renset data"]

    assert worksheet["B2"].value == 1250
    assert worksheet["C2"].value == 12500
    assert worksheet["D2"].value == 13000
    assert worksheet["E2"].value == 500
    assert worksheet["B2"].number_format == "#,##0"
    assert worksheet["C2"].number_format == "#,##0"
    assert worksheet["D2"].number_format == "#,##0"
    assert worksheet["E2"].number_format == "#,##0"


def test_build_export_workbook_formats_trailing_numeric_text_columns_even_with_generic_headers():
    cleaned_df = pd.DataFrame(
        [
            {
                "Ansvar": "100",
                "januar": "1 250,00",
                "Sumkolonne A": "12 500",
                "Sumkolonne B": "13 000",
                "Sumkolonne C": "500",
            }
        ]
    )

    workbook_bytes = build_export_workbook(
        cleaned_df=cleaned_df,
        adjustments_df=pd.DataFrame(columns=["Kategori", "Beskrivelse", "Verdi"]),
        user_notes="",
        source_name="pivotgrid.xlsx - Ark1",
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    worksheet = workbook["Renset data"]

    assert worksheet["C2"].value == 12500
    assert worksheet["D2"].value == 13000
    assert worksheet["E2"].value == 500
    assert worksheet["C2"].number_format == "#,##0"
    assert worksheet["D2"].number_format == "#,##0"
    assert worksheet["E2"].number_format == "#,##0"


def test_build_export_workbook_converts_accounting_style_negative_totals():
    cleaned_df = pd.DataFrame(
        [
            {
                "Ansvar": "100",
                "januar": "1 250,00",
                "Sumkolonne A": "(12 500)",
                "Sumkolonne B": "13 000-",
                "Sumkolonne C": "-500",
            }
        ]
    )

    workbook_bytes = build_export_workbook(
        cleaned_df=cleaned_df,
        adjustments_df=pd.DataFrame(columns=["Kategori", "Beskrivelse", "Verdi"]),
        user_notes="",
        source_name="pivotgrid.xlsx - Ark1",
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    worksheet = workbook["Renset data"]

    assert worksheet["C2"].value == -12500
    assert worksheet["D2"].value == -13000
    assert worksheet["E2"].value == -500
    assert worksheet["C2"].number_format == "#,##0"
    assert worksheet["D2"].number_format == "#,##0"
    assert worksheet["E2"].number_format == "#,##0"


def test_build_export_workbook_bolds_subtotal_rows():
    cleaned_df = pd.DataFrame(
        [
            {"Ktoniv7": "1001", "januar": 100.0, "Regnskap": 1000.0},
            {"Ktoniv7": "1002", "januar": 200.0, "Regnskap": 2000.0},
            {"Ktoniv7": "Lønnskostnader", "januar": 300.0, "Regnskap": 3000.0},
            {"Ktoniv7": "Ktoniv7", "januar": 300.0, "Regnskap": 3000.0},
        ]
    )
    cleaned_df.attrs["subtotal_labels"] = ["Lønnskostnader"]
    cleaned_df.attrs["total_label"] = "Ktoniv7"

    workbook_bytes = build_export_workbook(
        cleaned_df=cleaned_df,
        adjustments_df=pd.DataFrame(columns=["Kategori", "Beskrivelse", "Verdi"]),
        user_notes="",
        source_name="pivotgrid.xlsx - Ark1",
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    worksheet = workbook["Renset data"]

    assert worksheet["A4"].value == "Lønnskostnader"
    assert worksheet["A4"].font.bold is True
    assert worksheet["B4"].font.bold is True
    assert worksheet["C4"].font.bold is True
    assert worksheet["B4"].number_format == "#,##0"
    assert worksheet["C4"].number_format == "#,##0"
    assert worksheet["A5"].font.bold is True
    assert worksheet["B5"].font.bold is True
    assert worksheet["C5"].font.bold is True
